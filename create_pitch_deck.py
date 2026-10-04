from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
import pandas as pd


# =========================================================
# PITCHFORGE - STEP 5
# AUTOMATED PITCH DECK BUILDER
# =========================================================

print("==========================================")
print("       PITCHFORGE PITCH DECK BUILDER")
print("==========================================")


# =========================================================
# 1. LOAD DATA FROM STEP 3
# =========================================================

problem_df = pd.read_csv("problem_data.csv")
solution_df = pd.read_csv("solution_data.csv")
market_df = pd.read_csv("market_data.csv")
financial_df = pd.read_csv("financials_data.csv")

problem_text = " ".join(
    problem_df["Example Entry"].astype(str).tolist()
)

solution_text = " ".join(
    solution_df["Example Entry"].astype(str).tolist()
)


# =========================================================
# 2. CREATE PRESENTATION
# =========================================================

prs = Presentation()

# Widescreen format
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


# =========================================================
# 3. HELPER FUNCTIONS
# =========================================================

def add_title(slide, title):
    """Add slide title."""

    title_box = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(0.4),
        Inches(11.9),
        Inches(0.7)
    )

    text_frame = title_box.text_frame
    text_frame.clear()

    paragraph = text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(30)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(31, 78, 121)


def add_text(slide, text, x, y, width, height,
             font_size=20, bold=False):

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    text_frame = box.text_frame
    text_frame.word_wrap = True
    text_frame.clear()

    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(50, 50, 50)

    return box


def add_bullet(slide, text, x, y, width, height):

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    text_frame = box.text_frame
    text_frame.word_wrap = True
    text_frame.clear()

    paragraph = text_frame.paragraphs[0]
    paragraph.text = "• " + text
    paragraph.font.size = Pt(19)
    paragraph.font.color.rgb = RGBColor(50, 50, 50)

    return box


def add_footer(slide):
    """Add project status footer."""

    add_text(
        slide,
        "PitchForge | Proposed Startup Concept | Student Project",
        0.7,
        7.05,
        11.8,
        0.25,
        font_size=9
    )


# =========================================================
# SLIDE 1 - COVER
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_text(
    slide,
    "PITCHFORGE",
    0.8,
    1.8,
    11.7,
    0.8,
    font_size=42,
    bold=True
)

add_text(
    slide,
    "Build Your Story. Showcase Your Venture. Pitch with Confidence.",
    0.85,
    2.7,
    11.5,
    0.8,
    font_size=24
)

add_text(
    slide,
    "Digital Entrepreneur Portfolio & Pitch Deck Builder",
    0.85,
    3.7,
    11.5,
    0.6,
    font_size=20
)

add_text(
    slide,
    "Proposed Startup Concept",
    0.85,
    5.2,
    11.5,
    0.5,
    font_size=16,
    bold=True
)

add_footer(slide)


# =========================================================
# SLIDE 2 - PROBLEM
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_title(slide, "The Problem")

add_text(
    slide,
    "Early-stage entrepreneurs struggle to present their startup "
    "information in a structured, professional and investor-ready format.",
    0.8,
    1.5,
    11.5,
    1.5,
    font_size=25,
    bold=True
)

add_bullet(
    slide,
    "Information is often scattered across documents and presentations.",
    1.0,
    3.3,
    11,
    0.6
)

add_bullet(
    slide,
    "Many entrepreneurs lack professional presentation-design skills.",
    1.0,
    4.1,
    11,
    0.6
)

add_bullet(
    slide,
    "Creating a structured portfolio and pitch deck can be time-consuming.",
    1.0,
    4.9,
    11,
    0.6
)

add_footer(slide)


# =========================================================
# SLIDE 3 - SOLUTION
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_title(slide, "The Solution")

add_text(
    slide,
    "PitchForge",
    0.8,
    1.5,
    11.5,
    0.7,
    font_size=32,
    bold=True
)

add_text(
    slide,
    "An AI-assisted digital platform combining entrepreneur "
    "portfolio creation, startup information management and "
    "pitch-deck building in one platform.",
    0.8,
    2.3,
    11.5,
    1.5,
    font_size=23
)

add_bullet(
    slide,
    "Create a professional entrepreneur profile.",
    1.0,
    4.2,
    11,
    0.5
)

add_bullet(
    slide,
    "Organize startup information in one workspace.",
    1.0,
    4.9,
    11,
    0.5
)

add_bullet(
    slide,
    "Generate structured pitch-deck content.",
    1.0,
    5.6,
    11,
    0.5
)

add_footer(slide)


# =========================================================
# SLIDE 4 - MARKET
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_title(slide, "Target Market")

add_text(
    slide,
    "Primary Customer",
    0.8,
    1.4,
    5.5,
    0.5,
    font_size=25,
    bold=True
)

add_text(
    slide,
    "Student entrepreneurs and early-stage founders",
    0.8,
    2.0,
    5.5,
    1.0,
    font_size=21
)

add_text(
    slide,
    "Secondary Customers",
    6.8,
    1.4,
    5.5,
    0.5,
    font_size=25,
    bold=True
)

add_text(
    slide,
    "Incubators, accelerators, entrepreneurship cells "
    "and small-business owners",
    6.8,
    2.0,
    5.5,
    1.2,
    font_size=21
)

add_text(
    slide,
    "Early Adopter",
    0.8,
    3.8,
    5.5,
    0.5,
    font_size=25,
    bold=True
)

add_text(
    slide,
    "College students preparing for startup competitions, "
    "incubator applications or investor presentations.",
    0.8,
    4.4,
    11.2,
    1.2,
    font_size=21
)

add_text(
    slide,
    "Market-size figures: To be validated through market research.",
    0.8,
    6.0,
    11.5,
    0.5,
    font_size=14
)

add_footer(slide)


# =========================================================
# SLIDE 5 - BUSINESS MODEL
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_title(slide, "Business Model")

add_text(
    slide,
    "Proposed Revenue Model",
    0.8,
    1.3,
    11.5,
    0.6,
    font_size=27,
    bold=True
)

add_bullet(
    slide,
    "Freemium SaaS model",
    1.0,
    2.1,
    11,
    0.5
)

add_bullet(
    slide,
    "Paid subscription plans for advanced features",
    1.0,
    2.9,
    11,
    0.5
)

add_bullet(
    slide,
    "Institutional plans for colleges and entrepreneurship programmes",
    1.0,
    3.7,
    11,
    0.7
)

add_text(
    slide,
    "Proposed Pricing Structure",
    0.8,
    4.8,
    5.5,
    0.5,
    font_size=23,
    bold=True
)

add_text(
    slide,
    "Free → Basic features\n"
    "Pro → Full portfolio + AI assistance + pitch deck\n"
    "Premium → Advanced customization\n"
    "Institutional → Multiple accounts",
    0.8,
    5.4,
    11.5,
    1.4,
    font_size=17
)

add_footer(slide)


# =========================================================
# SLIDE 6 - TEAM
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_title(slide, "Team")

add_text(
    slide,
    "Founder",
    0.8,
    1.5,
    11.5,
    0.6,
    font_size=28,
    bold=True
)

add_text(
    slide,
    "Swetaa",
    0.8,
    2.2,
    11.5,
    0.6,
    font_size=25,
    bold=True
)

add_bullet(
    slide,
    "Student entrepreneur",
    1.0,
    3.2,
    11,
    0.5
)

add_bullet(
    slide,
    "Interest in entrepreneurship, digital business and AI tools",
    1.0,
    4.0,
    11,
    0.7
)

add_bullet(
    slide,
    "Responsible for startup concept, business strategy and project development",
    1.0,
    4.9,
    11,
    0.8
)

add_text(
    slide,
    "Team structure can be expanded as the startup progresses.",
    0.8,
    6.1,
    11.5,
    0.5,
    font_size=15
)

add_footer(slide)


# =========================================================
# SLIDE 7 - FINANCIALS
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_title(slide, "Financials")

add_text(
    slide,
    "Current Stage: Pre-Launch",
    0.8,
    1.4,
    11.5,
    0.6,
    font_size=27,
    bold=True
)

add_bullet(
    slide,
    "Revenue: Not yet generated",
    1.0,
    2.4,
    11,
    0.5
)

add_bullet(
    slide,
    "Customers: Not yet acquired",
    1.0,
    3.2,
    11,
    0.5
)

add_bullet(
    slide,
    "Financial projections: To be developed after validating pricing and cost assumptions",
    1.0,
    4.0,
    11,
    0.9
)

add_bullet(
    slide,
    "Funding requirement: To be estimated based on the next meaningful milestone",
    1.0,
    5.1,
    11,
    0.9
)

add_text(
    slide,
    "No actual revenue, customers or traction are claimed.",
    0.8,
    6.3,
    11.5,
    0.5,
    font_size=15,
    bold=True
)

add_footer(slide)


# =========================================================
# 8. SAVE PRESENTATION
# =========================================================

output_file = "PitchForge_Automated_Pitch_Deck.pptx"

prs.save(output_file)

print("\n==========================================")
print("PITCH DECK CREATED SUCCESSFULLY!")
print("==========================================")
print("\nCreated file:")
print(output_file)

print("\nSlides created:")
print("1. Cover")
print("2. Problem")
print("3. Solution")
print("4. Market")
print("5. Business Model")
print("6. Team")
print("7. Financials")