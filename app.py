import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader

from io import BytesIO


# =========================================================
# PITCHFORGE - STEP 6
# STREAMLIT WEB APPLICATION
# =========================================================

st.set_page_config(
    page_title="PitchForge",
    page_icon="🚀",
    layout="wide"
)


# =========================================================
# HEADER
# =========================================================

st.title("🚀 PitchForge")
st.subheader(
    "Digital Entrepreneur Portfolio & Pitch Deck Builder"
)

st.write(
    "Create your startup profile, receive smart content suggestions, "
    "upload visual assets and export a professional pitch deck."
)

st.divider()


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.title("PitchForge")

st.sidebar.info(
    "Prototype version\n\n"
    "Create → Improve → Export"
)


# =========================================================
# STARTUP INFORMATION
# =========================================================

st.header("1️⃣ Startup Information")

col1, col2 = st.columns(2)

with col1:

    startup_name = st.text_input(
        "Startup Name",
        value="PitchForge"
    )

    tagline = st.text_input(
        "Tagline",
        value="Build Your Story. Showcase Your Venture. Pitch with Confidence."
    )

    founder_name = st.text_input(
        "Founder Name",
        value="Swetaa"
    )

    industry = st.text_input(
        "Industry",
        value="Entrepreneurship Technology / SaaS / AI"
    )

with col2:

    target_market = st.text_area(
        "Target Market",
        value=(
            "Student entrepreneurs and early-stage founders. "
            "Secondary customers include incubators, accelerators "
            "and entrepreneurship cells."
        )
    )

    revenue_model = st.text_area(
        "Revenue Model",
        value=(
            "Freemium SaaS + paid subscriptions + "
            "institutional plans"
        )
    )


# =========================================================
# PROBLEM & SOLUTION
# =========================================================

st.header("2️⃣ Problem & Solution")

problem = st.text_area(
    "Problem Statement",
    value=(
        "Early-stage entrepreneurs struggle to organize and "
        "present their startup information in a professional "
        "and investor-ready format."
    ),
    height=120
)

solution = st.text_area(
    "Solution",
    value=(
        "PitchForge combines entrepreneur portfolio creation, "
        "startup information management and pitch-deck building "
        "in one digital platform."
    ),
    height=120
)


# =========================================================
# PRODUCT INFORMATION
# =========================================================

st.header("3️⃣ Product Information")

product = st.text_area(
    "Product / Service Description",
    value=(
        "A digital platform that helps entrepreneurs create "
        "professional portfolios and structured investor pitch decks."
    ),
    height=100
)

features = st.text_area(
    "Key Features",
    value=(
        "Founder Profile\n"
        "Portfolio Builder\n"
        "Startup Profile\n"
        "Pitch Deck Builder\n"
        "Smart Content Suggestions\n"
        "Export to PPTX and PDF"
    ),
    height=150
)


# =========================================================
# VISUAL UPLOADS
# =========================================================

st.header("4️⃣ Upload Visual Assets")

col1, col2, col3 = st.columns(3)

with col1:

    logo_file = st.file_uploader(
        "Upload Startup Logo",
        type=["png", "jpg", "jpeg"],
        key="logo"
    )

with col2:

    team_file = st.file_uploader(
        "Upload Team Photo",
        type=["png", "jpg", "jpeg"],
        key="team"
    )

with col3:

    product_file = st.file_uploader(
        "Upload Product Screenshot",
        type=["png", "jpg", "jpeg"],
        key="product"
    )


# =========================================================
# SMART CONTENT SUGGESTIONS
# =========================================================

st.header("5️⃣ Smart Content Suggestions")

if st.button("✨ Generate Suggestions"):

    if problem.strip():

        suggested_problem = (
            "Early-stage entrepreneurs struggle to communicate "
            "their startup opportunity in a structured and "
            "investor-ready format."
        )

    else:

        suggested_problem = (
            "Define the specific customer problem your startup solves."
        )

    if solution.strip():

        suggested_solution = (
            f"{startup_name} provides a structured digital platform "
            "that helps entrepreneurs transform startup information "
            "into professional portfolio and pitch-deck content."
        )

    else:

        suggested_solution = (
            "Describe how your product directly solves the identified problem."
        )

    st.success("Suggestions generated!")

    st.markdown("### Suggested Problem Statement")

    st.write(suggested_problem)

    st.markdown("### Suggested Solution Statement")

    st.write(suggested_solution)


# =========================================================
# PREVIEW
# =========================================================

st.header("6️⃣ Startup Preview")

preview_col1, preview_col2 = st.columns(2)

with preview_col1:

    st.markdown("### Startup")

    st.write(startup_name)

    st.markdown("### Tagline")

    st.write(tagline)

    st.markdown("### Founder")

    st.write(founder_name)

    st.markdown("### Industry")

    st.write(industry)

with preview_col2:

    st.markdown("### Problem")

    st.write(problem)

    st.markdown("### Solution")

    st.write(solution)

    st.markdown("### Target Market")

    st.write(target_market)


# =========================================================
# PPTX GENERATION FUNCTION
# =========================================================

def create_pptx():

    prs = Presentation()

    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)


    def add_title(slide, title):

        box = slide.shapes.add_textbox(
            Inches(0.7),
            Inches(0.4),
            Inches(11.8),
            Inches(0.7)
        )

        paragraph = box.text_frame.paragraphs[0]

        paragraph.text = title

        paragraph.font.size = Pt(30)
        paragraph.font.bold = True

        paragraph.font.color.rgb = RGBColor(
            31, 78, 121
        )


    def add_text(
        slide,
        text,
        x,
        y,
        width,
        height,
        size=20,
        bold=False
    ):

        box = slide.shapes.add_textbox(
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height)
        )

        tf = box.text_frame
        tf.word_wrap = True

        paragraph = tf.paragraphs[0]

        paragraph.text = text

        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold

        paragraph.font.color.rgb = RGBColor(
            50, 50, 50
        )


    # -----------------------------------------------------
    # COVER
    # -----------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_text(
        slide,
        startup_name,
        0.8,
        1.5,
        11.7,
        0.8,
        size=42,
        bold=True
    )

    add_text(
        slide,
        tagline,
        0.8,
        2.5,
        11.5,
        0.8,
        size=24
    )

    add_text(
        slide,
        "Digital Entrepreneur Portfolio & Pitch Deck",
        0.8,
        3.6,
        11.5,
        0.6,
        size=20
    )

    if logo_file:

        slide.shapes.add_picture(
            logo_file,
            Inches(10.7),
            Inches(0.8),
            width=Inches(1.6)
        )


    # -----------------------------------------------------
    # PROBLEM
    # -----------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(slide, "Problem")

    add_text(
        slide,
        problem,
        0.9,
        1.6,
        11.3,
        2.0,
        size=25,
        bold=True
    )

    add_text(
        slide,
        "Key Pain Points",
        0.9,
        4.0,
        11,
        0.5,
        size=23,
        bold=True
    )

    add_text(
        slide,
        "• Scattered startup information\n"
        "• Lack of structured templates\n"
        "• Difficulty creating professional presentations",
        1.0,
        4.7,
        11,
        1.5,
        size=19
    )


    # -----------------------------------------------------
    # SOLUTION
    # -----------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(slide, "Solution")

    add_text(
        slide,
        solution,
        0.9,
        1.6,
        11.3,
        1.8,
        size=24,
        bold=True
    )

    add_text(
        slide,
        product,
        0.9,
        3.8,
        11.3,
        1.0,
        size=20
    )

    add_text(
        slide,
        "Key Features",
        0.9,
        5.0,
        11,
        0.5,
        size=22,
        bold=True
    )

    add_text(
        slide,
        features,
        1.0,
        5.5,
        11,
        1.3,
        size=17
    )


    # -----------------------------------------------------
    # MARKET
    # -----------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(slide, "Target Market")

    add_text(
        slide,
        target_market,
        0.9,
        1.7,
        11.3,
        2.0,
        size=24
    )

    add_text(
        slide,
        "Market Position",
        0.9,
        4.2,
        11,
        0.5,
        size=23,
        bold=True
    )

    add_text(
        slide,
        "Focused on entrepreneurs who need both personal "
        "branding and startup pitch communication.",
        0.9,
        4.9,
        11,
        1.2,
        size=20
    )


    # -----------------------------------------------------
    # BUSINESS MODEL
    # -----------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(slide, "Business Model")

    add_text(
        slide,
        revenue_model,
        0.9,
        1.7,
        11.3,
        1.5,
        size=24,
        bold=True
    )

    add_text(
        slide,
        "Proposed Revenue Streams",
        0.9,
        3.8,
        11,
        0.5,
        size=23,
        bold=True
    )

    add_text(
        slide,
        "• Freemium access\n"
        "• Paid subscriptions\n"
        "• Premium features\n"
        "• Institutional plans",
        1.0,
        4.5,
        11,
        1.8,
        size=19
    )


    # -----------------------------------------------------
    # TEAM
    # -----------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(slide, "Founder / Team")

    add_text(
        slide,
        founder_name,
        0.9,
        1.6,
        7,
        0.7,
        size=30,
        bold=True
    )

    add_text(
        slide,
        "Founder – PitchForge",
        0.9,
        2.4,
        7,
        0.5,
        size=20
    )

    if team_file:

        slide.shapes.add_picture(
            team_file,
            Inches(8.5),
            Inches(1.4),
            width=Inches(3.5)
        )

    else:

        add_text(
            slide,
            "Team photo can be uploaded from the web app.",
            0.9,
            3.5,
            7,
            0.8,
            size=18
        )


    # -----------------------------------------------------
    # FINANCIALS
    # -----------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(slide, "Financial Overview")

    add_text(
        slide,
        "Current Stage: Pre-Launch",
        0.9,
        1.5,
        11,
        0.6,
        size=27,
        bold=True
    )

    add_text(
        slide,
        "Revenue: Not yet generated\n\n"
        "Customers: Not yet acquired\n\n"
        "Financial projections: To be developed after validating "
        "pricing and cost assumptions.\n\n"
        "Funding requirement: To be estimated based on the next "
        "meaningful milestone.",
        0.9,
        2.5,
        11.2,
        3.5,
        size=20
    )


    # -----------------------------------------------------
    # PRODUCT SCREENSHOT
    # -----------------------------------------------------

    if product_file:

        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        add_title(
            slide,
            "Product / MVP Preview"
        )

        slide.shapes.add_picture(
            product_file,
            Inches(1.2),
            Inches(1.5),
            width=Inches(10.8)
        )


    # -----------------------------------------------------
    # SAVE TO MEMORY
    # -----------------------------------------------------

    output = BytesIO()

    prs.save(output)

    output.seek(0)

    return output


# =========================================================
# PDF GENERATION FUNCTION
# =========================================================

def create_pdf():

    output = BytesIO()

    doc = SimpleDocTemplate(
        output,
        pagesize=A4
    )

    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph(
            startup_name,
            styles["Title"]
        )
    )

    story.append(
        Spacer(1, 15)
    )

    story.append(
        Paragraph(
            tagline,
            styles["Heading2"]
        )
    )

    story.append(
        Spacer(1, 20)
    )

    sections = [
        ("Founder", founder_name),
        ("Industry", industry),
        ("Problem", problem),
        ("Solution", solution),
        ("Product", product),
        ("Target Market", target_market),
        ("Revenue Model", revenue_model),
        ("Key Features", features)
    ]

    for title, content in sections:

        story.append(
            Paragraph(
                title,
                styles["Heading2"]
            )
        )

        story.append(
            Paragraph(
                content.replace("\n", "<br/>"),
                styles["BodyText"]
            )
        )

        story.append(
            Spacer(1, 12)
        )


    doc.build(story)

    output.seek(0)

    return output


# =========================================================
# EXPORT SECTION
# =========================================================

st.divider()

st.header("7️⃣ Export Your Pitch Materials")

col1, col2 = st.columns(2)

with col1:

    if st.button(
        "📊 Generate Pitch Deck",
        use_container_width=True
    ):

        pptx_file = create_pptx()

        st.session_state["pptx_file"] = (
            pptx_file.getvalue()
        )

        st.success(
            "Pitch deck generated successfully!"
        )


with col2:

    if st.button(
        "📄 Generate PDF",
        use_container_width=True
    ):

        pdf_file = create_pdf()

        st.session_state["pdf_file"] = (
            pdf_file.getvalue()
        )

        st.success(
            "PDF generated successfully!"
        )


# =========================================================
# DOWNLOAD BUTTONS
# =========================================================

if "pptx_file" in st.session_state:

    st.download_button(
        label="⬇️ Download Pitch Deck (.pptx)",
        data=st.session_state["pptx_file"],
        file_name="PitchForge_Pitch_Deck.pptx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
        use_container_width=True
    )


if "pdf_file" in st.session_state:

    st.download_button(
        label="⬇️ Download Portfolio (.pdf)",
        data=st.session_state["pdf_file"],
        file_name="PitchForge_Entrepreneur_Portfolio.pdf",
        mime="application/pdf",
        use_container_width=True
    )


# =========================================================
# FOOTER
# =========================================================

st.divider()

st.caption(
    "PitchForge | Digital Entrepreneur Portfolio & Pitch Deck Builder "
    "| Proposed Startup Concept"
)